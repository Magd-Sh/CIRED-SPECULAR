# -*- coding: utf-8 -*-
"""Point Hebdo 12 — style calqué sur Point_Hebdo_9 (blanc, sobre, texte noir)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image
import os

BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x55, 0x55)
FONT = "Calibri"
PLOTS = "/sessions/zealous-eager-dijkstra/mnt/data completion/imputation/outputs/plots"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def txt(s, x, y, w, h, runs, size=15, bold=False, color=BLACK, align=PP_ALIGN.LEFT,
        space_after=9, line=1.0):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    if isinstance(runs, str):
        runs = [(runs, {})]
    first = True
    for text, o in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = o.get("align", align)
        p.space_after = Pt(o.get("space_after", space_after))
        p.line_spacing = line
        content = ("•  " + text) if o.get("bullet") else text
        r = p.add_run()
        r.text = content
        f = r.font
        f.name = FONT
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = o.get("color", color)
    return tb

def title(s, t):
    txt(s, 0.7, 0.35, 11.9, 0.9, t, size=29)

def img(s, path, x, y, w=None, h=None):
    iw, ih = Image.open(path).size
    ar = iw / ih
    if w and not h:
        h = w / ar
    elif h and not w:
        w = h * ar
    s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h

# ---------------------------------------------------------------- 1 · Titre
s = slide()
txt(s, 1.2, 2.6, 10.9, 1.1, "Point Hebdo 12", size=48, align=PP_ALIGN.CENTER)
txt(s, 1.2, 3.8, 10.9, 1.2,
    "Imputation multiple de l’emploi : modèle, résultats, apport sur l’interpolation\n"
    "et chemin vers la question de recherche (taille des véhicules → emploi)",
    size=19, align=PP_ALIGN.CENTER, color=GREY)

# ------------------------------------------------- 2 · Les étapes
s = slide()
title(s, "Les étapes du pipeline (reproductible, Python)")
txt(s, 0.7, 1.35, 12.2, 5.6, [
    ("1)  Panel usine × année : 26 usines d’ASSEMBLAGE × 20 ans (2006–2025), emploi en logs.", {"space_after": 8}),
    ("2)  Covariables dérivées par usine-année, à partir de la BDD modèles/production :", {"space_after": 3}),
    ("taille = masse à vide moyenne des modèles actifs cette année-là (par usine) ; segment = moyenne d’un score ordinal Economy(1)…Luxury(6) ; part électrique = fraction de modèles 100 % électriques ; nombre de modèles = comptage à partir de la matrice de présence modèle × année.",
     {"size": 13, "color": GREY, "space_after": 8}),
    ("3)  Validation par masquage : on cache du connu, on impute, on compare — pour CHOISIR la méthode (cf. Hebdo 11).", {"space_after": 8}),
    ("4)  Imputation multiple : M = 10 panels complétés ; les cellules observées sont identiques, seules les imputées varient.", {"space_after": 8}),
    ("5)  Livrables : 10 panels + un résumé (estimation + intervalle + statut) + un utilitaire de mise en commun (règles de Rubin) pour les régressions.", {"space_after": 8}),
    ("6)  Graphiques de contrôle : trajectoires par usine et faisceaux de tirages (voir plus loin).", {}),
], size=14.5)

# ------------------------------------------------- 3 · Le modèle
s = slide()
title(s, "Le modèle retenu — et pourquoi pas les méthodes « lourdes »")
txt(s, 0.7, 1.35, 12.0, 1.5, [
    ("Constat de validation : sur CE panel, la complétion de matrice et l’imputation jointe (Amelia MVN) font MOINS bien que l’interpolation intra-usine. Peu de structure entre usines ; le signal dominant est la persistance propre à chaque usine.", {}),
], size=15)
tbl = s.shapes.add_table(4, 3, Inches(0.7), Inches(3.0), Inches(7.4), Inches(1.9)).table
tbl.columns[0].width = Inches(3.6); tbl.columns[1].width = Inches(1.9); tbl.columns[2].width = Inches(1.9)
rows = [
    ["Méthode", "Trous internes", "Trous de bord"],
    ["Interpolation intra-usine", "0,18", "0,19"],
    ["Amelia MVN jointe", "0,30", "0,47"],
    ["Imputation « within-plant » (retenue)", "0,18", "0,20"],
]
for i, row in enumerate(rows):
    for j, cell in enumerate(row):
        c = tbl.cell(i, j); c.text = cell
        for p in c.text_frame.paragraphs:
            for r in p.runs:
                r.font.name = FONT; r.font.size = Pt(12.5); r.font.bold = (i == 0)
                r.font.color.rgb = BLACK
        c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2) if i == 0 else RGBColor(0xFF, 0xFF, 0xFF)
txt(s, 8.4, 3.0, 4.3, 3.2, [
    ("La méthode retenue :", {"bold": True, "size": 15, "space_after": 5}),
    ("trous internes → interpolation du log-emploi + bruit ;", {"bullet": True, "size": 14}),
    ("trous de bord → marche ancrée sur la dernière obs., incréments = effet-année commun + variations de covariables, bruit croissant avec la distance ;", {"bullet": True, "size": 14}),
    ("M = 10 tirages → incertitude propagée (Rubin).", {"bullet": True, "size": 14}),
], size=14)
txt(s, 0.7, 5.1, 7.4, 0.5, "RMSE (log) — moyenne sur deux designs de masquage.", size=11, color=GREY)

# ------------------------------------------------- 4 · Les variables
s = slide()
title(s, "Les variables utilisées")
txt(s, 0.7, 1.35, 6.0, 5.4, [
    ("À imputer conjointement", {"bold": True, "size": 16, "space_after": 5}),
    ("log_emp — effectif (la cible)", {"bullet": True}),
    ("log_prod — production véhicules (elle-même trouée)", {"bullet": True, "space_after": 14}),
    ("Covariables (informent l’imputation)", {"bold": True, "size": 16, "space_after": 5}),
    ("kerb_weight — masse à vide moyenne = proxy de TAILLE des véhicules", {"bullet": True}),
    ("seg_score — segment moyen (ordinal)", {"bullet": True}),
    ("ev_share — part de modèles électriques", {"bullet": True}),
    ("n_models — nombre de modèles actifs", {"bullet": True}),
], size=14)
txt(s, 7.0, 1.35, 5.7, 5.5, [
    ("Structure (toujours observée)", {"bold": True, "size": 15, "space_after": 4}),
    ("effets fixes usine + tendance temporelle", {"bullet": True, "size": 13.5}),
    ("effet-année commun (chocs 2009, 2020…)", {"bullet": True, "size": 13.5, "space_after": 10}),
    ("Qualité des sources — un poids par valeur observée", {"bold": True, "size": 15, "space_after": 4}),
    ("Le code part de 1,0 et retranche selon la source : −0,15 si wiki (chiffre arrondi « ca. 20 000 »), −0,10 par année de report (carried_years) ; une source officielle (base / presse établie) reste à 1,0. Le poids est borné à [0,30 ; 1,0].",
     {"bullet": True, "size": 13.5}),
    ("Ce poids sert à l’over-imputation : les valeurs peu fiables sont traitées comme bruitées (moins de confiance), pas comme exactes.", {"bullet": True, "size": 13.5, "space_after": 10}),
    ("Taille véhicules : attributs ADEME propagés ; fallback nominal documenté pour 10 usines → couverture 41 % → 80 % (valeurs à vérifier).", {"bullet": True, "size": 13.5}),
], size=13.5)

# ------------------------------------------------- 5 · Usines omises
s = slide()
title(s, "Les usines écartées — et pourquoi")
txt(s, 0.7, 1.4, 12.0, 5.4, [
    ("Sites hors assemblage (14) : 12 usines de composants + 2 historiques → écartées. Elles ne produisent pas de véhicules, donc pas de « taille » — la variable centrale de l’étude. Périmètre final : 26 usines d’assemblage.", {"bullet": True}),
    ("Années « vides légitimes » (inactive) : hors de la plage d’activité d’une usine (usine pas encore ouverte / reconvertie) → NON imputées, pour ne pas inventer d’emploi là où l’usine n’existait pas.", {"bullet": True}),
    ("Usines sous 25 % de couverture emploi (Porsche Leipzig, Września, Osnabrück) : imputées mais FLAGGÉES « faible couverture » — à traiter avec prudence ou à compléter par recherche presse, plutôt que de laisser une série imputée à 80 %+ peser dans les régressions.", {"bullet": True}),
    ("Point d’attention : le statut « inactive » dépend de la BDD modèles ; une usine active tôt mais sans donnée modèle (ex. Emden avant 2010) est marquée inactive → règle ajustable si une usine est exclue à tort.", {"bullet": True}),
], size=15)

# ------------------------------------------------- 6 · Apport vs interpolation (+ grille)
s = slide()
title(s, "Apport sur l’interpolation simple")
txt(s, 0.7, 1.25, 5.0, 5.6, [
    ("Les prédictions ponctuelles sont ~équivalentes (l’imputation reproduit l’interpolation sur les trous internes). L’apport est ailleurs :", {"space_after": 12}),
    ("Incertitude quantifiée : 10 tirages → bande / écart-type par cellule → écarts-types honnêtes via Rubin. L’interpolation seule donne une ligne à fausse confiance.", {"bullet": True}),
    ("Covariables sur les bords : la marche ancrée utilise production, taille, effet-année là où l’interpolation ne fait qu’un plat.", {"bullet": True}),
    ("Imputation conjointe et cohérente de l’emploi et de la production.", {"bullet": True}),
    ("Traçabilité : chaque cellule flaggée (observée / interne / bord / faible couverture).", {"bullet": True}),
], size=14)
img(s, os.path.join(PLOTS, "employment_grid.png"), 6.0, 1.15, w=6.9)
txt(s, 6.0, 6.75, 6.9, 0.4,
    "26 usines — pointillé = interpolation (avant) · bleu + bande = imputation (après).",
    size=10.5, color=GREY, align=PP_ALIGN.CENTER)

# ------------------------------------------------- 7 · L'incertitude (spaghetti)
s = slide()
title(s, "L’incertitude visualisée — les 10 tirages, par groupe")
img(s, os.path.join(PLOTS, "spaghetti_VW.png"), 0.7, 1.25, w=6.6)
img(s, os.path.join(PLOTS, "spaghetti_MERCEDES.png"), 7.5, 1.25, w=5.5)
img(s, os.path.join(PLOTS, "spaghetti_BMW.png"), 7.5, 3.55, w=5.5)
txt(s, 0.7, 6.55, 12.0, 0.7, [
    ("Les 10 lignes se resserrent sur chaque point observé (certitude) et s’écartent sur les bords / longs trous (incertitude) — c’est exactement ce que Rubin fait entrer dans les écarts-types.", {"size": 12.5, "color": GREY}),
], size=12.5)

# ------------------------------------------------- 8 · Question de recherche + endogénéité
s = slide()
title(s, "Vers la question de recherche : taille des véhicules → emploi")
txt(s, 0.7, 1.25, 12.2, 5.7, [
    ("Oui, c’est estimable : régression emploi ~ taille véhicule + contrôles (effets fixes usine + année) lancée sur CHACUN des 10 panels, puis combinée par les règles de Rubin. L’écart-type final intègre l’incertitude d’imputation.", {"space_after": 10}),
    ("Éviter l’endogénéité — trois points :", {"bold": True, "size": 15.5, "space_after": 4}),
    ("Pas de circularité : garder la variable de taille DANS le modèle d’imputation (déjà le cas). L’imputation ne fabrique pas la relation — elle prolonge celle des cellules observées, et Rubin en paie le prix (écarts-types plus larges). L’omettre atténuerait le coefficient vers zéro.", {"bullet": True, "size": 13.5}),
    ("Vrai risque = missingness non-aléatoire (la presse couvre surtout les crises / plans sociaux) : tests de robustesse — régression sur cellules observées seules, en excluant les usines très imputées, et une imputation « sans la taille » comme borne inférieure.", {"bullet": True, "size": 13.5}),
    ("Production = médiateur, pas un simple contrôle : réduire la taille des voitures agit probablement SUR l’emploi en partie À TRAVERS la production (plateformes partagées, moins d’heures de montage par véhicule…). Si on met la production à droite de la régression, on absorbe ce canal et on n’estime plus que l’effet « à production donnée » — souvent proche de zéro, alors même que l’effet total existe. Décision à prendre : effet TOTAL (ne pas contrôler la production, mais la garder dans l’imputation) ou effet DIRECT (la contrôler), et l’assumer explicitement. Comme aucune identification causale n’est revendiquée, on présentera cela comme des associations, la production servant surtout de covariable d’imputation.", {"bullet": True, "size": 13.5}),
], size=14)

# ------------------------------------------------- 9 · Prochaines étapes
s = slide()
title(s, "Limites et prochaines étapes")
txt(s, 0.7, 1.35, 12.2, 5.6, [
    ("Bords pré-2010 : irréductiblement incertains (~30 % d’erreur en test) — flaggés, intervalles larges. Envisager de borner l’analyse à 2010–2025.", {"bullet": True, "size": 14.5}),
    ("Production trouée (37 % observée) : limite l’apport des covariables là où emploi ET production manquent (premières années).", {"bullet": True, "size": 14.5}),
    ("Fallback taille à raffiner : pour 10 usines, la masse des véhicules vient de valeurs constructeur nominales (par famille de modèle), pas de specs vérifiées. Or la taille est LA variable explicative : une mesure bruitée atténue son coefficient (biais vers zéro). Priorité = remplacer ces valeurs par des specs sourcées (homologation UE, extension ADEME) — cela améliore à la fois l’imputation et la régression finale.", {"bullet": True, "size": 14.5}),
    ("Prochain jalon = la régression elle-même : emploi ~ taille + contrôles (effets fixes usine + année) sur chacun des 10 panels, puis mise en commun par Rubin (l’utilitaire est déjà écrit). On y ajoutera d’emblée les tests de robustesse (cellules observées seules, exclusion des usines très imputées, borne « sans la taille »), puis on décidera du périmètre temporel (2010–2025 ?) au vu des résultats.", {"bullet": True, "size": 14.5}),
], size=14.5)

out = "/sessions/zealous-eager-dijkstra/mnt/data completion/Point_Hebdo_12.pptx"
prs.save(out)
print("saved", out)
